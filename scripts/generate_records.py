"""
generate_records.py – Generate synthetic employee records full of (fake) PII for demos.

Every value is random: names are drawn from common Indian first names and surnames, and the
PAN, Aadhaar and phone numbers only match the formats, not any real person.

    python scripts/generate_records.py                          # 1000 records -> employee_records_1000.txt
    python scripts/generate_records.py --count 100 --format docx
"""
import argparse
import random
from pathlib import Path

FIRST_NAMES = [
    "Arjun", "Aditi", "Rohan", "Priya", "Vikram", "Sneha", "Karan", "Ananya", "Rahul", "Ishita",
    "Siddharth", "Kavita", "Amit", "Meera", "Varun", "Tanvi", "Akash", "Riya", "Manish", "Shreya",
    "Deepak", "Neha", "Sanjay", "Pooja", "Rajesh", "Aarti", "Vijay", "Jyoti", "Sunil", "Sunita",
    "Anil", "Anita", "Suresh", "Seema", "Ramesh", "Rekha", "Dinesh", "Kiran", "Mukesh", "Lata",
    "Pankaj", "Maya", "Abhishek", "Ritu", "Vivek", "Sonal", "Harish", "Swati", "Nitin", "Madhu",
    "Pradeep", "Rashmi", "Gaurav", "Shweta", "Sandeep", "Babita", "Ashok", "Suman", "Arvind", "Bindu",
    "Yogesh", "Geeta", "Kuldeep", "Manju", "Satish", "Chitra", "Ajay", "Deepa", "Vinod", "Preeti",
    "Naveen", "Vandana", "Alok", "Mamta", "Manoj", "Sarita", "Hemant", "Nisha", "Shailesh", "Sangeeta",
    "Bhavesh", "Komal", "Dharmendra", "Sushma", "Ganesh", "Laxmi", "Tushar", "Poonam", "Lalit", "Shanti",
    "Paresh", "Vimla", "Kishore", "Usha", "Prashant", "Kusum", "Subhash", "Nirmala", "Jagdish", "Savita",
]

SURNAMES = [
    "Sharma", "Verma", "Gupta", "Malhotra", "Mehra", "Kapoor", "Khanna", "Joshi", "Patel", "Shah",
    "Singh", "Yadav", "Trivedi", "Desai", "Kulkarni", "Patil", "Reddy", "Nair", "Iyer", "Rao",
    "Agarwal", "Bansal", "Goel", "Chopra", "Das", "Chatterjee", "Banerjee", "Mukherjee", "Bose", "Ghosh",
    "Pandey", "Mishra", "Tiwari", "Dubey", "Shukla", "Pathak", "Chaudhary", "Dhillon", "Gill", "Sandhu",
]

CITIES = [
    "Mumbai", "Delhi", "Bengaluru", "Hyderabad", "Ahmedabad", "Chennai", "Kolkata", "Surat", "Pune",
    "Jaipur", "Lucknow", "Kanpur", "Nagpur", "Indore", "Thane", "Bhopal", "Visakhapatnam",
    "Pimpri-Chinchwad", "Patna", "Vadodara",
]

DEPARTMENTS = ["Engineering", "Human Resources", "Marketing", "Sales", "Finance", "Legal"]
LETTERS = "ABCDEFGHIJKLMNOPQRSTUVWXYZ"


def make_record(rng: random.Random, i: int) -> str:
    first, last = rng.choice(FIRST_NAMES), rng.choice(SURNAMES)
    pan = "".join(rng.choice(LETTERS) for _ in range(5)) + str(rng.randint(1000, 9999)) + rng.choice(LETTERS)
    street = rng.choice(["Garden", "Sunrise", "Palm", "Ocean"])
    return "\n".join([
        f"EMPLOYEE RECORD #{1000 + i}",
        f"Name: {first} {last}",
        f"DOB: {rng.randint(1, 28)}-{rng.randint(1, 12)}-{rng.randint(1975, 2000)}",
        f"Email: {first.lower()}.{last.lower()}{i}@example.in",
        f"Phone: +91-{rng.randint(70000, 99999)}-{rng.randint(10000, 99999)}",
        f"PAN Card: {pan}",
        f"Aadhaar: {rng.randint(2000, 9999)}-{rng.randint(1000, 9999)}-{rng.randint(1000, 9999)}",
        f"Address: Flat {rng.randint(101, 999)}, {street} View, {rng.choice(CITIES)}, "
        f"Maharashtra {rng.randint(400001, 400999)}",
        f"Work Department: {rng.choice(DEPARTMENTS)}",
    ])


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--count", type=int, default=1000)
    parser.add_argument("--format", choices=["txt", "docx"], default="txt")
    parser.add_argument("--seed", type=int, default=7, help="same seed, same records")
    parser.add_argument("--out", type=Path, help="defaults to employee_records_<count>.<format>")
    args = parser.parse_args()

    rng = random.Random(args.seed)
    records = [make_record(rng, i) for i in range(1, args.count + 1)]
    out = args.out or Path(f"employee_records_{args.count}.{args.format}")

    if args.format == "txt":
        out.write_text("\n\n---\n\n".join(records), encoding="utf-8")
    else:
        from docx import Document

        doc = Document()
        for record in records:
            for line in record.splitlines():
                doc.add_paragraph(line)
            doc.add_paragraph("---")
        doc.save(out)
    print(f"Wrote {len(records)} records to {out}")


if __name__ == "__main__":
    main()
